from __future__ import annotations

from functools import lru_cache

from openpyxl import load_workbook

from .config import settings
from .schemas import HazardLabel


@lru_cache(maxsize=1)
def load_catalog() -> list[HazardLabel]:
    if not settings.catalog_excel.exists():
        raise FileNotFoundError(f"标签Excel不存在：{settings.catalog_excel}")
    workbook = load_workbook(settings.catalog_excel, read_only=True, data_only=True)
    if settings.catalog_sheet not in workbook.sheetnames:
        raise ValueError(f"工作表不存在：{settings.catalog_sheet}")
    sheet = workbook[settings.catalog_sheet]
    labels: list[HazardLabel] = []
    end_row = settings.catalog_first_row + settings.catalog_label_count - 1
    try:
        for row in sheet.iter_rows(min_row=settings.catalog_first_row, max_row=end_row, values_only=True):
            sequence, category, normalized_name, inspection_method = row[0], row[1], row[4], row[8]
            if sequence is None or not normalized_name:
                continue
            label_id = f"H{int(sequence):03d}"
            labels.append(
                HazardLabel(
                    id=label_id,
                    category=str(category or "未分类").strip(),
                    name=str(normalized_name).strip(),
                    inspection_method=str(inspection_method or "").strip(),
                )
            )
    finally:
        workbook.close()
    if len(labels) != settings.catalog_label_count:
        raise ValueError(f"应读取{settings.catalog_label_count}个标签，实际读取{len(labels)}个")
    return labels


@lru_cache(maxsize=1)
def catalog_by_id() -> dict[str, HazardLabel]:
    return {item.id: item for item in load_catalog()}


def labels_for_categories(categories: set[str], limit: int | None = None) -> list[HazardLabel]:
    matched = [item for item in load_catalog() if item.category in categories]
    if not matched:
        matched = load_catalog()
    return matched[:limit] if limit else matched

